# -*- coding: utf-8 -*-
"""
Chuyen doi cac file docx tuan 1-35 cua Lop 1 va Lop 2 thanh js/khbd_sohoa/lopX/lopX_tin_hoc.js
chuan dinh dang web thuviengiaoanbaigiang
"""
import os
import sys
import json
import docx

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

def parse_docx_week(file_path, week_num):
    doc = docx.Document(file_path)
    
    lesson_title = f"TUẦN {week_num}"
    topic = ""
    yccd = []
    dodung = []
    activities = []
    dieuchinh = []
    tables_data = []
    
    current_section = None
    
    # Duyệt paragraphs
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
            
        t_upper = t.upper()
        
        # Tiêu đề tuần / chủ đề / bài
        if t_upper.startswith("TUẦN ") and len(t) < 30:
            continue
        elif "CHỦ ĐỀ" in t_upper and len(t) < 100 and not topic:
            topic = t
            continue
        elif (t_upper.startswith("BÀI ") or "TÊN BÀI HỌC:" in t_upper) and len(t) < 120 and lesson_title == f"TUẦN {week_num}":
            lesson_title = t.replace("Tên bài học:", "").strip()
            continue
            
        # Section detection
        if "A. YÊU CẦU CẦN ĐẠT" in t_upper or "I. YÊU CẦU CẦN ĐẠT" in t_upper or "A.  YÊU CẦU CẦN ĐẠT" in t_upper:
            current_section = "yccd"
            yccd.append(t)
            continue
        elif "B. THIẾT BỊ DẠY HỌC" in t_upper or "II. ĐỒ DÙNG DẠY HỌC" in t_upper or "PHƯƠNG TIỆN, THIẾT BỊ DẠY HỌC" in t_upper or "II. THIẾT BỊ" in t_upper:
            current_section = "dodung"
            dodung.append(t)
            continue
        elif "C. CÁC HOẠT ĐỘNG" in t_upper or "III. CÁC HOẠT ĐỘNG" in t_upper or "C. TIẾN TRÌNH DẠY HỌC" in t_upper:
            current_section = "activities"
            continue
        elif "D. ĐIỀU CHỈNH" in t_upper or "IV. ĐIỀU CHỈNH" in t_upper:
            current_section = "dieuchinh"
            continue
            
        # Append lines
        if current_section == "yccd":
            yccd.append(t)
        elif current_section == "dodung":
            dodung.append(t)
        elif current_section == "activities":
            activities.append(t)
        elif current_section == "dieuchinh":
            dieuchinh.append(t)
            
    # Duyệt tables
    for tbl in doc.tables:
        rows_data = []
        for r in tbl.rows:
            row_cells = []
            for c in r.cells:
                # Lấy text từng cell
                cell_text = "\n".join([p.text.strip() for p in c.paragraphs if p.text.strip()])
                row_cells.append(cell_text)
            rows_data.append(row_cells)
        if rows_data:
            tables_data.append(rows_data)
            
    if not dodung:
        dodung = [
            "- Đối với GV: Máy tính, máy chiếu, bài trình chiếu PowerPoint, tệp thực hành mẫu.",
            "- Đối với HS: Sách giáo khoa, vở ghi chép, máy vi tính."
        ]
        
    lesson_obj = {
        "lessonTitle": lesson_title,
        "topic": topic,
        "yccd": yccd,
        "dodung": dodung,
        "activities": activities,
        "dieuchinh": dieuchinh,
        "tables": tables_data
    }
    
    return {
        "week": week_num,
        "sourceFile": f"TUẦN {week_num}.docx",
        "lessons": [lesson_obj]
    }

def generate_js_for_grade(grade):
    src_dir = rf"D:\Giao án mới\TINHOC_CHUAN_HOA\Lop {grade}"
    out_js = rf"c:\Users\Admin\Desktop\Thư viện Giáo án - Bài giảng\js\khbd_sohoa\lop{grade}\lop{grade}_tin_hoc.js"
    os.makedirs(os.path.dirname(out_js), exist_ok=True)
    
    weeks_data = {}
    for w in range(1, 36):
        docx_file = os.path.join(src_dir, f"TUẦN {w}.docx")
        if os.path.exists(docx_file):
            weeks_data[str(w)] = parse_docx_week(docx_file, w)
        else:
            print(f"Warning: Missing TUAN {w}.docx for Grade {grade}")
            
    js_content = f"""(function() {{
  var grade = {grade};
  var subjectId = "tin_hoc";
  var subjectName = "Tin học";
  var weeksData = {json.dumps(weeks_data, ensure_ascii=False, indent=2)};

  if (typeof KHBD_DATA !== 'undefined' && KHBD_DATA.registerSubject) {{
    KHBD_DATA.registerSubject(grade, subjectId, subjectName, weeksData);
  }}
}})();
"""
    with open(out_js, "w", encoding="utf-8") as f:
        f.write(js_content)
        
    print(f"Đã tạo thành công {out_js} với {len(weeks_data)} tuần!")

if __name__ == '__main__':
    generate_js_for_grade(1)
    generate_js_for_grade(2)
